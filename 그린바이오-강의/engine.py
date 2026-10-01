# ===== 그린바이오 전자상거래개론 강의자료 생성 엔진 (python-pptx + python-docx) =====
import re, copy
from pptx import Presentation
from pptx.oxml.ns import qn

R_NS = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}"
A = "{http://schemas.openxmlformats.org/drawingml/2006/main}"

# 슬롯 이름 → (도형 index, 문단 index)  ※ 5주차_1차시 원본 기준, 절대 바꾸지 말 것
SLOTS = {
 1: {"photo":(2,0),"week":(4,0),"title1":(5,0),"title2":(5,1),"tagline":(7,0)},
 2: {"subtitle":(2,0),"c1_title":(5,0),"c1_l1":(7,0),"c1_l2":(7,1),
     "c2_title":(10,0),"c2_l1":(12,0),"c2_l2":(12,1),
     "c3_title":(15,0),"c3_l1":(17,0),"c3_l2":(17,1),"bottom":(19,0)},
 3: {"photo":(2,0),"title1":(6,0),"title2":(6,1),"sub1":(8,0),"sub2":(8,1)},
 4: {"label":(0,0),"title":(1,0),"subtitle":(2,0),"r1_h":(5,0),"r1_d":(6,0),"r2_h":(9,0),"r2_d":(10,0),
     "r3_h":(13,0),"r3_d":(14,0),"r4_h":(17,0),"r4_d":(18,0)},
 5: {"label":(0,0),"title":(1,0),"subtitle":(2,0),
     "a_period":(4,0),"a_value":(5,0),"a_label":(6,0),"a_d1":(8,0),"a_d2":(8,1),
     "b_period":(10,0),"b_value":(11,0),"b_label":(12,0),"b_d1":(14,0),"b_d2":(14,1),"source":(15,0)},
 6: {"label":(0,0),"title":(1,0),"subtitle":(2,0),
     "k1_title":(5,0),"k1_l1":(7,0),"k1_l2":(7,1),"k2_title":(10,0),"k2_l1":(12,0),"k2_l2":(12,1),
     "k3_title":(15,0),"k3_l1":(17,0),"k3_l2":(17,1),"k4_title":(20,0),"k4_l1":(22,0),"k4_l2":(22,1),"bottom":(24,0)},
 7: {"photo":(2,0),"title1":(6,0),"title2":(6,1),"sub1":(8,0),"sub2":(8,1)},
 8: {"label":(0,0),"title":(1,0),"subtitle":(2,0),"r1_h":(5,0),"r1_d":(6,0),"r2_h":(9,0),"r2_d":(10,0),
     "r3_h":(13,0),"r3_d":(14,0),"r4_h":(17,0),"r4_d":(18,0)},
 9: {"label":(0,0),"title":(1,0),"subtitle":(2,0),"r1_h":(5,0),"r1_d":(6,0),"r2_h":(9,0),"r2_d":(10,0),
     "r3_h":(13,0),"r3_d":(14,0),"r4_h":(17,0),"r4_d":(18,0)},
 10:{"title1":(2,0),"title2":(2,1),"body1":(3,0),"body2":(3,1)},
 11:{"title":(1,0),"subtitle":(2,0),"r1_h":(5,0),"r1_d":(6,0),"r2_h":(9,0),"r2_d":(10,0),
     "r3_h":(13,0),"r3_d":(14,0),"next_label":(16,0),"next_title":(17,0)},
}
ROADMAP_CARDS = [3, 8, 13]   # 2번 장표 카드 사각형 index (1·2·3차시)

def _set(shape, pidx, text):
    ps = [p for p in shape.text_frame.paragraphs if "".join(r.text for r in p.runs).strip()]
    if pidx >= len(ps): return False
    ps[pidx].runs[0].text = text
    for r in ps[pidx].runs[1:]: r.text = ""
    return True

def _line(shape, w, color):
    ln = shape._element.find(".//" + A + "ln")
    ln.set("w", str(w))
    ln.find(A + "solidFill").find(A + "srgbClr").set("val", color)

def build(template, out, week, cha, content, notes, foot_old="WEEK 5 · 1차시"):
    """content = {"S1.title1": "...", ...}, notes = [11개 대본 문자열]"""
    assert len(notes) == 11, "대본은 11개여야 합니다"
    prs = Presentation(template); miss = []
    for si, slide in enumerate(prs.slides, 1):
        shapes = list(slide.shapes)
        for name, (idx, p) in SLOTS[si].items():
            key = f"S{si}.{name}"
            if key in content and not _set(shapes[idx], p, content[key]): miss.append(key)
        # 푸터
        for sh in shapes:
            if sh.has_text_frame:
                for para in sh.text_frame.paragraphs:
                    for r in para.runs:
                        if foot_old in r.text: r.text = r.text.replace(foot_old, f"WEEK {week} · {cha}차시")
        # 로드맵 강조 카드 이동
        if si == 2:
            for k, ci in enumerate(ROADMAP_CARDS, 1):
                _line(shapes[ci], 13970 if k == cha else 9525, "14A178" if k == cha else "065640")
        # 원본 사진(그림 N)·미니맥스 아이콘 제거 (+관계 정리)
        for sh in shapes:
            if sh.shape_type == 13 and (sh.name.startswith("그림") or sh.name.startswith("MiniMax")):
                el = sh._element
                rids = {v for n in el.iter() for k, v in n.attrib.items() if k.startswith(R_NS)}
                el.getparent().remove(el)
                for rid in rids:
                    used = any(v == rid for n in slide._element.iter() for k, v in n.attrib.items() if k.startswith(R_NS))
                    if not used and rid in slide.part.rels: slide.part.drop_rel(rid)
        # 음성 자동재생 설정 제거 (남으면 파워포인트가 파일을 못 엶)
        for t in slide._element.findall(qn("p:timing")): slide._element.remove(t)
        slide.notes_slide.notes_text_frame.text = notes[si - 1]
    prs.save(out)
    return miss

def qa(path, old_words=("WEEK 5",)):
    prs = Presentation(path); rep = {}
    notes = [s.notes_slide.notes_text_frame.text for s in prs.slides]
    allnote = " ".join(notes)
    rep["슬라이드 수"] = len(prs.slides._sldIdLst)
    rep["대본 글자수"] = len(allnote)
    rep["예상 분량(분)"] = round(len(allnote) / 449, 1)
    rep["대본 속 아라비아 숫자"] = sorted(set(re.findall(r"[0-9]+", allnote))) or "없음"
    bad = 0
    for s in prs.slides:
        x = s._element.xml
        bad += len(set(re.findall(r'spid="(\d+)"', x)) - set(re.findall(r'<p:cNvPr id="(\d+)"', x)))
        bad += 100 if "<p:timing" in x else 0
    rep["끊긴참조/자동재생 잔여"] = bad
    rep["이전 텍스트 잔여"] = [(i, w) for i, s in enumerate(prs.slides, 1) for sh in s.shapes
                          if sh.has_text_frame for w in old_words if w in sh.text_frame.text] or "없음"
    rep["슬라이드별 대본 글자수"] = [len(n) for n in notes]
    return rep

def script_docx(path, course_line, title_line, slide_titles, notes):
    from docx import Document
    from docx.shared import Pt, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml.ns import qn as dqn
    d = Document()
    st = d.styles["Normal"]; st.font.name = "Malgun Gothic"; st.font.size = Pt(10.5)
    st.element.rPr.rFonts.set(dqn("w:eastAsia"), "Malgun Gothic")
    total = sum(len(n) for n in notes)
    for txt, size, color in [(course_line, 10, "5F6A5E"), (title_line, 17, "1B3A2C"),
                             (f"미니맥스 TTS 녹음용 (11슬라이드) — 숫자 한글화 · {total:,}자 · 약 {total/449:.1f}분", 9.5, "5F6A5E")]:
        p = d.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(txt); r.bold = size > 10; r.font.size = Pt(size); r.font.color.rgb = RGBColor.from_string(color)
    for i, (t, n) in enumerate(zip(slide_titles, notes), 1):
        h = d.add_paragraph(); r = h.add_run(f"슬라이드 {i}. {t}"); r.bold = True
        r.font.size = Pt(12); r.font.color.rgb = RGBColor.from_string("0F7A5C")
        for para in [x for x in n.split("\n") if x.strip()]:
            d.add_paragraph(para.strip())
    d.save(path)
